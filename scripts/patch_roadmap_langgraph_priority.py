#!/usr/bin/env python3
"""Elevate LangGraph to Phase 0 — the first roadmap priority, ahead of every
existing item — by surgically editing the existing Roadmap slides in place
(no script ever survived in this repo that rebuilds slides 10-12 from
scratch, so this edits their existing shapes/paragraphs directly instead of
re-drawing the slide). Also repositions the 3 LangGraph slides to sit right
after "Production layout" and before "Roadmap — trust & quality", so the
LangGraph story reads first in the roadmap narrative. Idempotent.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs" / "deliverables" / "v0.1" / "zooplus-assistant-interview-15min-pro.pptx"

TITLE_INK = RGBColor(0x0A, 0x25, 0x40)
BODY_INK = RGBColor(0x0F, 0x17, 0x2A)

PHASE0_LINES = [
    ("Phase 0 — DONE: LangGraph state-graph foundation", True),
    ("feature/lg-catalog-probe-intent — faithful redesign, audited live (see previous slides)", False),
]

LG_TITLES = (
    "LangGraph 1/3",
    "LangGraph 2/3",
    "LangGraph 3/3",
)


def _find_slide_by_title(prs: Presentation, title_prefix: str) -> int | None:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and (sh.text or "").strip().startswith(title_prefix):
                return i
    return None


def _find_shape_with_text(slide, text_prefix: str):
    for sh in slide.shapes:
        if sh.has_text_frame and (sh.text or "").strip().startswith(text_prefix):
            return sh
    return None


def _rebuild_bullets(shape, lines: list[tuple[str, bool]], *, size_pt: float) -> None:
    tf = shape.text_frame
    # Clear all existing paragraphs except the first (python-pptx requires
    # at least one paragraph to remain); reuse it for the new first line.
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    first.clear()
    run = first.add_run()
    run.text = lines[0][0]
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = lines[0][1]
    run.font.color.rgb = TITLE_INK if lines[0][1] else BODY_INK
    first.space_after = Pt(8)
    for text, bold in lines[1:]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.name = "Arial"
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.color.rgb = TITLE_INK if bold else BODY_INK
        p.space_after = Pt(8)


def _update_phase0_slide(prs: Presentation) -> None:
    idx = _find_slide_by_title(prs, "Roadmap — trust")
    if idx is None:
        print("WARNING: 'Roadmap — trust & quality' slide not found, skipping")
        return
    slide = prs.slides[idx]

    subtitle_shape = _find_shape_with_text(slide, "Phase 1")
    if subtitle_shape is not None:
        subtitle_shape.text_frame.paragraphs[0].text = (
            "Phase 0 done · Phase 1 P0 — LangGraph foundation, then security (items 0, 2–3)"
        )

    body_shape = _find_shape_with_text(slide, "Why first")
    if body_shape is None:
        print("WARNING: roadmap body bullets not found on 'Roadmap — trust' slide")
        return
    existing_lines = [(p.text, bool(p.font.bold)) for p in body_shape.text_frame.paragraphs]
    if existing_lines and existing_lines[0][0].startswith("Phase 0"):
        return  # already patched
    size_pt = body_shape.text_frame.paragraphs[0].font.size.pt
    _rebuild_bullets(body_shape, PHASE0_LINES + existing_lines, size_pt=size_pt)


def _update_priority_order(prs: Presentation) -> None:
    idx = _find_slide_by_title(prs, "Roadmap — product channels")
    if idx is None:
        print("WARNING: 'Roadmap — product channels' slide not found, skipping")
        return
    slide = prs.slides[idx]
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            if p.text.strip().startswith("Priority order:") and "0→2" not in p.text and p.runs:
                p.runs[0].text = "Priority order: 0→2→3→8→9→4→5→7→1→6"


def _move_slide_to(prs: Presentation, from_index: int, to_index: int) -> None:
    xml_slides = prs.slides._sldIdLst
    elems = list(xml_slides)
    elem = elems[from_index]
    xml_slides.remove(elem)
    xml_slides.insert(to_index, elem)


def _reposition_langgraph_slides(prs: Presentation) -> None:
    for title in LG_TITLES:
        src = _find_slide_by_title(prs, title)
        dst = _find_slide_by_title(prs, "Roadmap — trust")
        if src is None or dst is None:
            continue
        if src == dst:
            continue
        target = dst if src > dst else dst - 1
        if src == target:
            continue
        _move_slide_to(prs, src, target)


def _renumber_footer(prs: Presentation) -> None:
    from pptx.util import Inches

    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            if sh.top and sh.top > Inches(6.9) and sh.width and sh.width < Inches(1.2):
                t = sh.text_frame.text.strip()
                if t.isdigit():
                    sh.text_frame.paragraphs[0].text = str(i)


def main() -> None:
    if not DECK.is_file():
        raise SystemExit(f"Missing deck: {DECK}")
    prs = Presentation(str(DECK))

    _reposition_langgraph_slides(prs)
    _update_phase0_slide(prs)
    _update_priority_order(prs)
    _renumber_footer(prs)

    tmp = DECK.with_name(f"{DECK.stem}_p0{DECK.suffix}")
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, DECK)
        tmp.unlink(missing_ok=True)
        print(f"Updated {DECK} — LangGraph repositioned as Phase 0, priority order updated")
    except PermissionError:
        print(f"PPT is open — saved to {tmp}")
        print("Close PowerPoint, then copy the patched file over the deck.")


if __name__ == "__main__":
    main()
