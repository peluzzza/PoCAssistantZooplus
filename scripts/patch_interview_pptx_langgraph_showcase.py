#!/usr/bin/env python3
"""Insert a dedicated "LangGraph approach" slide before the closing
Release-progress slide. This deck (releases, v2.1.6 baseline) carries no
LangGraph code — the slide documents the approach and points to the
standalone showcase branch where it's actually implemented. Idempotent —
re-running updates content instead of duplicating the slide.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from patch_interview_pptx_rag_slides import (  # type: ignore[import-not-found]
    PANEL_BG,
    _bullets,
    draw_pro_chrome,
    insert_slide_at,
    renumber_footer,
)
from patch_interview_pptx_rag_slides import _rgb as _rgb_panel  # type: ignore[import-not-found]
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs" / "deliverables" / "v0.1" / "zooplus-assistant-interview-15min-pro.pptx"

SLIDE_TITLE = "LangGraph approach — standalone showcase branch"

LEFT_LG = [
    ("feature/lg-catalog-probe-intent", True),
    ("• Standalone branch — not merged into dev/main/releases.", False),
    ("• Same agentic core re-expressed as an explicit LangGraph StateGraph", False),
    ("  (nodes + typed conditional edges) instead of Python if/else.", False),
    ("• Design rule: every node calls THIS repo's own orchestrator", False),
    ("  functions directly — zero re-implemented logic, nothing invented.", False),
    ("Why explore a state graph", True),
    ("• Retry/fallback decisions become named, independently unit-tested", False),
    ("  routing functions instead of nested control flow in one function.", False),
    ("• Declarative node/edge graph — visualizable, replayable per request.", False),
]

RIGHT_LG = [
    ("Topology on that branch", True),
    ("• cache_check -> classify_intent -> [route_by_lane]", False),
    ("•      -> social (decline_off_topic | conversational)", False),
    ("•      -> prefetch -> process_lane (catalog_search)", False),
    ("• -> cache_store -> END", False),
    ("• process_lane calls dispatch_process(run_process_lane) directly —", False),
    ("  the same ACP dispatch this repo's catalog lane already used.", False),
    ("What a full audit found", True),
    ("• 2 real gaps in an earlier, looser translation: UI model override", False),
    ("  silently ignored; no chat-level cache. Both closed by calling the", False),
    ("  real functions instead of re-deriving equivalent-looking logic.", False),
]


def build_langgraph_slide(slide, *, page: int) -> None:
    draw_pro_chrome(
        slide,
        title=SLIDE_TITLE,
        subtitle="Faithful node/edge translation — every node calls this repo's own functions",
        badge="Architecture · LG",
        page=page,
    )
    panel = slide.shapes.add_shape(1, Inches(0.64), Inches(1.38), Inches(12.0), Inches(5.45))
    _rgb_panel(panel, PANEL_BG)
    _bullets(
        slide,
        Inches(0.82),
        Inches(1.55),
        Inches(5.85),
        Inches(5.1),
        LEFT_LG,
        size=13,
    )
    _bullets(
        slide,
        Inches(6.7),
        Inches(1.55),
        Inches(5.75),
        Inches(5.1),
        RIGHT_LG,
        size=13,
    )


def _find_slide_by_title(prs: Presentation, title_prefix: str) -> int | None:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and (sh.text or "").strip().startswith(title_prefix):
                return i
    return None


def _ensure_langgraph_slide(prs: Presentation) -> None:
    existing = _find_slide_by_title(prs, "LangGraph approach")
    progress_idx = _find_slide_by_title(prs, "Release progress")
    if existing is not None:
        slide = prs.slides[existing]
        for sh in list(slide.shapes):
            el = sh._element
            el.getparent().remove(el)
        build_langgraph_slide(slide, page=existing + 1)
        return
    insert_at = progress_idx if progress_idx is not None else len(prs.slides)
    slide = insert_slide_at(prs, insert_at)
    build_langgraph_slide(slide, page=insert_at + 1)


def main() -> None:
    if not DECK.is_file():
        raise SystemExit(f"Missing deck: {DECK}")
    prs = Presentation(str(DECK))

    _ensure_langgraph_slide(prs)
    renumber_footer(prs)

    tmp = DECK.with_name(f"{DECK.stem}_lg{DECK.suffix}")
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, DECK)
        tmp.unlink(missing_ok=True)
        print(f"Updated {DECK} — now {len(prs.slides)} slides (LangGraph approach slide added)")
    except PermissionError:
        print(f"PPT is open — saved to {tmp}")
        print("Close PowerPoint, then copy the patched file over the deck.")


if __name__ == "__main__":
    main()
