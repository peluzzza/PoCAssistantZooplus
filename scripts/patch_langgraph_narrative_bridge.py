#!/usr/bin/env python3
"""Fix narrative coherence around the LangGraph slides (run after
patch_roadmap_langgraph_priority.py has repositioned them between
"Production layout" and "Roadmap — trust & quality"). Three problems this
fixes:

1. Slide 9 -> 10 had zero transition: it jumped from "Production layout"
   straight to "The problem this explores" with no sentence connecting them.
2. Slide 12 said "standalone, not merged" (an exploration) while slide 13
   said "Phase 0 — DONE" (implying shipped/adopted) — contradictory tone.
3. Slide 13 claimed LangGraph "shapes how every other roadmap item gets
   implemented" without saying which items or how — an unsupported
   assertion a listener can't verify.

Idempotent — checks for its own markers before re-applying.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs" / "deliverables" / "v0.1" / "zooplus-assistant-interview-15min-pro.pptx"

TITLE_INK = RGBColor(0x0A, 0x25, 0x40)
BODY_INK = RGBColor(0x0F, 0x17, 0x2A)

BRIDGE_LINE = "Also explored for FR5: would an explicit state graph beat nested Python branches? (next 3 slides)"
CONCRETE_LINE = "If adopted: item 2 (intent facets) -> new conditional edge; item 5 (MCP/ACP bus) -> new node types"


def _find_slide_by_title(prs: Presentation, title_prefix: str) -> int | None:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and (sh.text or "").strip().startswith(title_prefix):
                return i
    return None


def _find_shape_with_text(slide, text_prefix: str):
    for sh in slide.shapes:
        if sh.has_text_frame and (sh.text or "").strip().startswith(text_prefix):
            return sh
    return None


def _set_paragraph_text(p, text: str) -> None:
    """Replace a paragraph's visible text while keeping its first run's
    formatting (size/bold/color) intact — paragraph.text would instead drop
    to default formatting on a fresh run."""
    if p.runs:
        p.runs[0].text = text
    else:
        p.text = text


def _append_paragraph(shape, text: str, *, bold: bool = False) -> None:
    tf = shape.text_frame
    ref = tf.paragraphs[0]
    size_pt = ref.font.size.pt if ref.font.size else 16
    space_after = ref.space_after
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = TITLE_INK if bold else BODY_INK
    if space_after is not None:
        p.space_after = space_after


def _bridge_production_slide(prs: Presentation) -> None:
    idx = _find_slide_by_title(prs, "Production layout")
    if idx is None:
        return
    slide = prs.slides[idx]
    body = _find_shape_with_text(slide, "FR5 = Production-oriented")
    if body is None:
        return
    if any(p.text.strip() == BRIDGE_LINE for p in body.text_frame.paragraphs):
        return  # already patched
    _append_paragraph(body, BRIDGE_LINE)


def _tag_langgraph_badges(prs: Presentation) -> None:
    """Append '· #0' to the 'Architecture · LG' badge on all 3 LangGraph
    slides so slide 13's "(items 0, 2-3)" reference is traceable back to a
    visible label, instead of an unexplained bare number."""
    for title in ("LangGraph 1/3", "LangGraph 2/3", "LangGraph 3/3"):
        idx = _find_slide_by_title(prs, title)
        if idx is None:
            continue
        slide = prs.slides[idx]
        badge = _find_shape_with_text(slide, "Architecture")
        if badge is None:
            continue
        p = badge.text_frame.paragraphs[0]
        if p.text.strip().endswith("#0"):
            continue
        _set_paragraph_text(p, "Architecture · LG · #0")


def _fix_closing_line(prs: Presentation) -> None:
    """Slide 12 (LangGraph 3/3): replace the flat 'standalone, not merged'
    line with one that still tells the truth (not merged) but also bridges
    forward into why it's positioned first in the roadmap next."""
    idx = _find_slide_by_title(prs, "LangGraph 3/3")
    if idx is None:
        return
    slide = prs.slides[idx]
    new_text = (
        "Explored & validated on feature/lg-catalog-probe-intent (not yet adopted) — "
        "recommended as item #0, ahead of the roadmap below."
    )
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            if p.text.strip().startswith("Branch: feature/lg-catalog-probe-intent — standalone"):
                _set_paragraph_text(p, new_text)


def _fix_phase0_subtitle(slide) -> None:
    subtitle = _find_shape_with_text(slide, "Phase 0")
    if subtitle is None:
        return
    new_text = "Item #0 explored · Phase 1 P0 next — LangGraph foundation, then security (items 0, 2–3)"
    _set_paragraph_text(subtitle.text_frame.paragraphs[0], new_text)


def _fix_phase0_bullets(slide) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        paragraphs = sh.text_frame.paragraphs
        if not any(p.text.strip().startswith("Phase 0 — DONE") for p in paragraphs):
            continue
        for p in paragraphs:
            if p.text.strip().startswith("Phase 0 — DONE"):
                _set_paragraph_text(p, "Item #0 — explored & validated (not yet adopted)")
        if not any(p.text.strip() == CONCRETE_LINE for p in paragraphs):
            _append_paragraph(sh, CONCRETE_LINE)
        return


def _fix_phase0_wording(prs: Presentation) -> None:
    """Slide 13 (Roadmap — trust & quality): soften 'Phase 0 — DONE' (which
    contradicts 'not merged' on the previous slide) to 'explored & validated',
    and add a concrete sentence explaining which later items this affects and
    how, instead of an unsupported 'shapes everything' claim."""
    idx = _find_slide_by_title(prs, "Roadmap — trust")
    if idx is None:
        return
    slide = prs.slides[idx]
    _fix_phase0_subtitle(slide)
    _fix_phase0_bullets(slide)


def main() -> None:
    if not DECK.is_file():
        raise SystemExit(f"Missing deck: {DECK}")
    prs = Presentation(str(DECK))

    _bridge_production_slide(prs)
    _tag_langgraph_badges(prs)
    _fix_closing_line(prs)
    _fix_phase0_wording(prs)

    tmp = DECK.with_name(f"{DECK.stem}_nb{DECK.suffix}")
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, DECK)
        tmp.unlink(missing_ok=True)
        print(f"Updated {DECK} — narrative bridge + concrete connections added")
    except PermissionError:
        print(f"PPT is open — saved to {tmp}")
        print("Close PowerPoint, then copy the patched file over the deck.")


if __name__ == "__main__":
    main()
